#!/usr/bin/env python3
"""Wrap the graphical-abstract figure in a one-slide PowerPoint with presenter notes.

Run scripts/build_graphical_abstract.py first; this places its PNG full-bleed on
a 16:9 slide and adds notes walking through the four equations.
Requires python-pptx.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]

NOTES = """\
OVERVIEW
We want the age T of an ancient sample from the alleles it carries. The idea: treat the ancient genome as one more lineage threaded into the ARG of the modern panel. Its age decides where that lineage can join the trees, and that decides which derived alleles it can carry. Read the slide left to right.

PANEL 1 - INSERT THE LINEAGE
The gray tree is one local tree of the modern ARG (8 tips here; 26 haplotypes in our data). The blue lineage starts at the sample's age T and moves back in time until it coalesces with ("joins") the tree at time u.
Equation: S_T(u) = exp[ - integral from T to u of kappa(s) / 2Ne(s) ds ].
- kappa(s) is the number of tree lineages present at time s. The ancient lineage coalesces with each one at rate 1/(2Ne), the usual coalescent rate, so its total rate of joining is kappa/2Ne.
- S_T(u) is the probability that it has NOT joined yet by time u: a survival function. It falls fast when the tree has many lineages (recent times) and slowly deep in the tree.
- u is never estimated. It is a latent variable that we integrate out in the next step. Ne(t) can vary through time; the integral just uses the local value.

PANEL 2 - ONE CURVE PER SITE
For each SNP i, the mutation sits on one edge of the tree (a, b, c on the left). The ancient sample carries the derived allele only if its lineage joins the tree below that mutation, i.e. on a branch descended from it, and before the mutation happened.
Equation: q_i(T) = integral from T to z_i of S_T(u) * delta_i(u) / 2Ne(u) du.
- delta_i(u) is the number of tree lineages at time u that descend from the mutation edge, so delta_i/2Ne is the rate of joining a "derived" lineage.
- In words: the probability of surviving unjoined to u, times the rate of joining a derived lineage at u, summed over every possible u.
- The upper limit z_i is the mutation's time. If the sample is older than the mutation (T > z_i), q = 0: an ancestor cannot carry a mutation that had not happened yet.
- We do not know z_i exactly, only its edge, so q is averaged uniformly over the edge's time span (MATH.md eq. 15). Between tree nodes everything is piecewise constant, so the integrals have closed forms and are fast.
- The curves are computed with the pipeline code on this exact tree. Near T = 0, q is higher for mutations carried by more of the modern tips (c is in 6 of 8 tips, b in 4, a in 2), though below that fraction because the lineage can also join above the mutation. q falls as T grows and reaches zero at the top of the edge. Young mutations are informative only about young samples.

PANEL 3 - COMBINE SITES
Each site is a coin flip: x_i = 1 if the sample shows the derived allele, which has probability q_i(T), and x_i = 0 for ancestral, probability 1 - q_i(T).
Equation: L(T) = product over sites of q_i(T)^x_i * [1 - q_i(T)]^(1 - x_i).
- That is just the Bernoulli likelihood multiplied across tens of thousands of SNPs. With a flat prior on T, the posterior P(T | data) is L(T) normalised.
- This example is one simulated sample, about 40,000 sites, analysed with the true modern ARG. True age about 5,100 generations; estimate 5,050; 95% interval about 4,560 to 5,600. The axis is zoomed so the width is visible. (Chosen as a clean illustration; see caveats.)
- Shown without error for simplicity. In real data, q is replaced by eps + (1 - 2 eps) q, where eps is an effective discordance absorbing genotype error and ARG error. With posterior ARG samples we multiply sites within each ARG draw, average over draws, and multiply chromosomes.
- The product treats sites as independent even though linked sites are not. It is a composite likelihood: the point estimate is fine, but the interval is too narrow.

PANEL 4 - SIMULATION CHECK
T-hat is the posterior mode (MAP). Each point is one simulated ancient sample, true age against estimate.
- Blue: 300 simulations with constant Ne (Ne, mutation and recombination rates, modern panel size and sample age varied between runs), no error. Mean error +13 generations, slope 1.00.
- Orange: 1,000 samples under varied, piecewise-constant Ne(t) histories with 1% genotype error. Mean error +8 generations, slope 0.99.
- Both use the true modern ARG. So the likelihood itself, including variable Ne and the error model, gives essentially unbiased ages.

CAVEATS (if asked)
- Intervals: because the site product ignores linkage, nominal 95% intervals cover the truth only about 46% of the time. A 5 Mb block bootstrap helps but is not yet nominal.
- Inferred ARGs (SINGER topology + POLEGON dating) introduce demography-dependent bias, and no single eps value removes it. That is the current focus.
"""


def build(image: Path, out: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    pic = slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width,
                                   height=prs.slide_height)
    pic.name = "Graphical abstract"
    pic._element.nvPicPr.cNvPr.set(
        "descr",
        "Four panels: an ancient lineage inserted into a modern tree; per-site "
        "derived-allele probability curves; the age posterior for one sample; "
        "estimated versus true age in simulations.")
    slide.notes_slide.notes_text_frame.text = NOTES
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(out)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--image", type=Path, default=ROOT / "docs/graphical_abstract_slide.png")
    p.add_argument("--out", type=Path, default=ROOT / "docs/graphical_abstract_slide.pptx")
    a = p.parse_args()
    build(a.image, a.out)


if __name__ == "__main__":
    main()
